import re
from datetime import datetime
from pathlib import Path
from typing import Callable, TypeVar

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.models.dae_models import (
    DaeParceladoEncontrado,
    DaeParceladoNaoEncontrado,
    DaePagamentoNaoLocalizado,
    LinhaIcmsAt,
    NotaConciliada,
    NotaNaoEncontrada,
    PagamentoConfirmado,
)

T = TypeVar("T")

COLUNAS_CONCILIADAS = [
    "Número da NF",
    "Código da Receita",
    "Referência",
    "Valor Principal (R$)",
    "Especificação da Receita",
    "Status",
]

COLUNAS_NAO_ENCONTRADAS = [
    "Número da NF",
    "Data de Emissão",
    "CNPJ do Emitente",
    "Status",
]

COLUNAS_PAGAMENTOS_CONFIRMADOS = [
    "Nosso Número",
    "Pagamento",
    "Referência",
    "Receita",
    "Val. Principal",
    "Val. Total",
]

COLUNAS_PAGAMENTOS_NAO_LOCALIZADOS = [
    "Arquivo de Origem",
    "Código da Receita",
    "Referência",
    "Valor Principal (R$)",
    "Status",
]

COLUNAS_PARCELAMENTO_ENCONTRADOS = [
    "Arquivo DAE",
    "PAF",
    "Data Ocorrência",
    "Data Vencimento",
    "Valor Histórico",
    "Valor Débito",
]

COLUNAS_PARCELAMENTO_NAO_ENCONTRADOS = [
    "Arquivo de Origem",
    "Código da Receita",
    "Referência",
    "Valor Principal (R$)",
    "Status",
]

COLUNAS_ICMS_AT = [
    "Valor Apurado (R$)",
    "Valor Pago (Código 1.145) (R$)",
    "Valor a Recolher (R$)",
]


PADRAO_REFERENCIA_COMPLETA = re.compile(r"^\s*(\d{1,2})/(\d{4})\s*$")
PADRAO_DATA_ISO = re.compile(r"^(\d{4})-(\d{2})-\d{2}")
PADRAO_DATA_BR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})")


def _ano_mes_da_referencia(referencia: str | None) -> tuple[int, int] | None:
    if not referencia:
        return None
    m = PADRAO_REFERENCIA_COMPLETA.match(referencia)
    return (int(m.group(2)), int(m.group(1))) if m else None


def _ano_mes_da_data_emissao(data_emissao: str | None) -> tuple[int, int] | None:
    if not data_emissao:
        return None
    texto = data_emissao.strip()
    m = PADRAO_DATA_ISO.match(texto)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = PADRAO_DATA_BR.match(texto)
    if m:
        return int(m.group(3)), int(m.group(2))
    return None


def _titulo_ano(ano: int | None) -> str:
    return f"Ano {ano}" if ano is not None else "Sem Ano Identificado"


def _titulo_mes(mes: int | None) -> str:
    return f"Mês {mes:02d}" if mes is not None else "Sem Mês Identificado"


def _agrupar_por_ano_mes(
    itens: list[T], ano_mes_de: Callable[[T], tuple[int, int] | None]
) -> list[tuple[int | None, list[tuple[int | None, list[T]]]]]:
    por_ano_mes: dict[int | None, dict[int | None, list[T]]] = {}
    for item in itens:
        resultado = ano_mes_de(item)
        ano, mes = resultado if resultado is not None else (None, None)
        por_ano_mes.setdefault(ano, {}).setdefault(mes, []).append(item)

    grupos = []
    for ano in sorted(por_ano_mes, key=lambda a: (a is None, a)):
        meses = por_ano_mes[ano]
        grupos.append((ano, [(mes, meses[mes]) for mes in sorted(meses, key=lambda m: (m is None, m))]))
    return grupos


def _escrever_planilha_por_ano_mes(
    planilha,
    colunas: list[str],
    grupos_ano: list[tuple[int | None, list[tuple[int | None, list[T]]]]],
    linha_de: Callable[[T], list],
) -> None:
    if not grupos_ano:
        planilha.append(colunas)
        return
    primeiro_ano = True
    for ano, grupos_mes in grupos_ano:
        if not primeiro_ano:
            planilha.append([])
        primeiro_ano = False
        planilha.append([_titulo_ano(ano)])
        primeiro_mes = True
        for mes, itens in grupos_mes:
            if not primeiro_mes:
                planilha.append([])
            primeiro_mes = False
            planilha.append([_titulo_mes(mes)])
            planilha.append(colunas)
            for item in itens:
                planilha.append(linha_de(item))


def gerar_relatorio_conciliadas(conciliadas: list[NotaConciliada], caminho: Path) -> Path:
    workbook = Workbook()
    planilha = workbook.active
    grupos = _agrupar_por_ano_mes(conciliadas, lambda n: _ano_mes_da_referencia(n.referencia))
    _escrever_planilha_por_ano_mes(
        planilha,
        COLUNAS_CONCILIADAS,
        grupos,
        lambda nota: [
            nota.numero_nf,
            nota.codigo_receita,
            nota.referencia,
            nota.valor_principal,
            nota.especificacao_receita,
            f"🟢 {nota.status}",
        ],
    )
    workbook.save(str(caminho))
    return caminho


def _elementos_por_ano_mes(
    estilos,
    colunas: list[str],
    grupos_ano: list[tuple[int | None, list[tuple[int | None, list[T]]]]],
    linha_de: Callable[[T], list],
    cor_cabecalho: str,
    estilo_titulo_ano: str = "Heading2",
    estilo_titulo_mes: str = "Heading3",
) -> list:
    if not grupos_ano:
        return [_tabela_relatorio([colunas], cor_cabecalho)]
    elementos = []
    for ano, grupos_mes in grupos_ano:
        elementos.append(Paragraph(_titulo_ano(ano), estilos[estilo_titulo_ano]))
        for mes, itens in grupos_mes:
            elementos.append(Paragraph(_titulo_mes(mes), estilos[estilo_titulo_mes]))
            dados = [colunas] + [linha_de(item) for item in itens]
            elementos.append(_tabela_relatorio(dados, cor_cabecalho))
            elementos.append(Spacer(1, 8))
    return elementos


def gerar_relatorio_conciliadas_pdf(conciliadas: list[NotaConciliada], caminho: Path) -> Path:
    documento = SimpleDocTemplate(
        str(caminho),
        pagesize=landscape(A4),
        title="AuditaDAE - Notas Conciliadas",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("AuditaDAE — Notas Conciliadas", estilos["Title"]),
        Paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Spacer(1, 12),
    ]

    grupos = _agrupar_por_ano_mes(conciliadas, lambda n: _ano_mes_da_referencia(n.referencia))
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_CONCILIADAS,
            grupos,
            lambda nota: [
                nota.numero_nf,
                nota.codigo_receita or "",
                nota.referencia or "",
                f"{nota.valor_principal:.2f}" if nota.valor_principal is not None else "",
                nota.especificacao_receita or "",
                nota.status,
            ],
            "#1f6f43",
            "Heading2",
            "Heading3",
        )
    )
    documento.build(elementos)
    return caminho


def gerar_relatorio_nao_encontradas_pdf(nao_encontradas: list[NotaNaoEncontrada], caminho: Path) -> Path:
    documento = SimpleDocTemplate(
        str(caminho),
        pagesize=landscape(A4),
        title="AuditaDAE - Notas Nao Encontradas",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("AuditaDAE — Notas Não Encontradas", estilos["Title"]),
        Paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Spacer(1, 12),
    ]

    grupos = _agrupar_por_ano_mes(nao_encontradas, lambda n: _ano_mes_da_data_emissao(n.data_emissao))
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_NAO_ENCONTRADAS,
            grupos,
            lambda nota: [
                nota.numero_nf,
                nota.data_emissao or "",
                nota.cnpj_emitente or "",
                nota.status,
            ],
            "#8f1f1f",
            "Heading2",
            "Heading3",
        )
    )
    documento.build(elementos)
    return caminho


def gerar_relatorio_nao_encontradas(nao_encontradas: list[NotaNaoEncontrada], caminho: Path) -> Path:
    workbook = Workbook()
    planilha = workbook.active
    grupos = _agrupar_por_ano_mes(nao_encontradas, lambda n: _ano_mes_da_data_emissao(n.data_emissao))
    _escrever_planilha_por_ano_mes(
        planilha,
        COLUNAS_NAO_ENCONTRADAS,
        grupos,
        lambda nota: [
            nota.numero_nf,
            nota.data_emissao,
            nota.cnpj_emitente,
            f"🔴 {nota.status}",
        ],
    )
    workbook.save(str(caminho))
    return caminho


def _linha_pagamento_confirmado_excel(confirmado: PagamentoConfirmado) -> list:
    return [
        confirmado.linha_pagamento.nosso_numero,
        confirmado.linha_pagamento.data_pagamento,
        confirmado.linha_pagamento.referencia_bruta,
        f"{confirmado.linha_pagamento.codigo_receita or ''} - {confirmado.linha_pagamento.descricao_receita or ''}",
        confirmado.linha_pagamento.valor_principal,
        confirmado.linha_pagamento.valor_total,
    ]


def _linha_pagamento_nao_localizado_excel(nao_localizado: DaePagamentoNaoLocalizado) -> list:
    return [
        nao_localizado.arquivo_origem,
        nao_localizado.codigo_receita,
        nao_localizado.referencia,
        nao_localizado.valor_principal,
        nao_localizado.status,
    ]


def _ano_mes_pagamento_confirmado(confirmado: PagamentoConfirmado) -> tuple[int, int] | None:
    ano, mes = confirmado.linha_pagamento.ano_referencia, confirmado.linha_pagamento.mes_referencia
    return (ano, mes) if ano is not None and mes is not None else None


def gerar_relatorio_pagamentos_excel(
    confirmados: list[PagamentoConfirmado],
    nao_localizados: list[DaePagamentoNaoLocalizado],
    caminho: Path,
) -> Path:
    workbook = Workbook()

    aba_confirmados = workbook.active
    aba_confirmados.title = "Confirmados"
    grupos_confirmados = _agrupar_por_ano_mes(confirmados, _ano_mes_pagamento_confirmado)
    _escrever_planilha_por_ano_mes(
        aba_confirmados, COLUNAS_PAGAMENTOS_CONFIRMADOS, grupos_confirmados, _linha_pagamento_confirmado_excel
    )

    aba_nao_localizados = workbook.create_sheet("Não Localizados")
    grupos_nao_localizados = _agrupar_por_ano_mes(nao_localizados, lambda n: _ano_mes_da_referencia(n.referencia))
    _escrever_planilha_por_ano_mes(
        aba_nao_localizados, COLUNAS_PAGAMENTOS_NAO_LOCALIZADOS, grupos_nao_localizados, _linha_pagamento_nao_localizado_excel
    )

    workbook.save(str(caminho))
    return caminho


def _tabela_relatorio(dados: list[list], cor_cabecalho: str) -> Table:
    tabela = Table(dados, repeatRows=1)
    tabela.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(cor_cabecalho)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return tabela


def gerar_relatorio_pagamentos_pdf(
    confirmados: list[PagamentoConfirmado],
    nao_localizados: list[DaePagamentoNaoLocalizado],
    caminho: Path,
) -> Path:
    documento = SimpleDocTemplate(
        str(caminho),
        pagesize=landscape(A4),
        title="AuditaDAE - Relatorio de Pagamentos",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("AuditaDAE — Relatório de Pagamentos", estilos["Title"]),
        Paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Spacer(1, 12),
    ]

    elementos.append(Paragraph("Pagamentos Confirmados", estilos["Heading2"]))
    grupos_confirmados = _agrupar_por_ano_mes(confirmados, _ano_mes_pagamento_confirmado)
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_PAGAMENTOS_CONFIRMADOS,
            grupos_confirmados,
            lambda confirmado: [
                confirmado.linha_pagamento.nosso_numero,
                confirmado.linha_pagamento.data_pagamento or "",
                confirmado.linha_pagamento.referencia_bruta or "",
                f"{confirmado.linha_pagamento.codigo_receita or ''} - {confirmado.linha_pagamento.descricao_receita or ''}",
                f"{confirmado.linha_pagamento.valor_principal:.2f}" if confirmado.linha_pagamento.valor_principal is not None else "",
                f"{confirmado.linha_pagamento.valor_total:.2f}" if confirmado.linha_pagamento.valor_total is not None else "",
            ],
            "#1f6f43",
            "Heading3",
            "Heading4",
        )
    )
    elementos.append(Spacer(1, 12))

    elementos.append(Paragraph("Não Localizados no Relatório de Pagamentos", estilos["Heading2"]))
    grupos_nao_localizados = _agrupar_por_ano_mes(nao_localizados, lambda n: _ano_mes_da_referencia(n.referencia))
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_PAGAMENTOS_NAO_LOCALIZADOS,
            grupos_nao_localizados,
            lambda nao_localizado: [
                nao_localizado.arquivo_origem,
                nao_localizado.codigo_receita or "",
                nao_localizado.referencia or "",
                f"{nao_localizado.valor_principal:.2f}" if nao_localizado.valor_principal is not None else "",
                nao_localizado.status,
            ],
            "#8f1f1f",
            "Heading3",
            "Heading4",
        )
    )

    documento.build(elementos)
    return caminho


def gerar_relatorio_icms_at_excel(linhas: list[LinhaIcmsAt], caminho: Path) -> Path:
    workbook = Workbook()
    planilha = workbook.active
    grupos = _agrupar_por_ano_mes(linhas, lambda l: (l.ano, l.mes))
    _escrever_planilha_por_ano_mes(
        planilha,
        COLUNAS_ICMS_AT,
        grupos,
        lambda l: [l.valor_apurado, l.valor_pago, l.valor_a_recolher],
    )
    workbook.save(str(caminho))
    return caminho


def gerar_relatorio_icms_at_pdf(linhas: list[LinhaIcmsAt], caminho: Path) -> Path:
    documento = SimpleDocTemplate(
        str(caminho),
        pagesize=landscape(A4),
        title="AuditaDAE - Buscas ICMS Antecipacao Tributaria",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("AuditaDAE — Buscas ICMS Antecipação Tributária", estilos["Title"]),
        Paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Spacer(1, 12),
    ]

    grupos = _agrupar_por_ano_mes(linhas, lambda l: (l.ano, l.mes))
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_ICMS_AT,
            grupos,
            lambda l: [
                f"{l.valor_apurado:.2f}" if l.valor_apurado is not None else "",
                f"{l.valor_pago:.2f}" if l.valor_pago is not None else "",
                f"{l.valor_a_recolher:.2f}" if l.valor_a_recolher is not None else "",
            ],
            "#1f4f8f",
            "Heading2",
            "Heading3",
        )
    )
    documento.build(elementos)
    return caminho


def _linha_parcelamento_encontrado_excel(encontrado: DaeParceladoEncontrado) -> list:
    return [
        encontrado.dae_arquivo_origem,
        encontrado.linha_parcelamento.paf,
        encontrado.linha_parcelamento.data_ocorrencia,
        encontrado.linha_parcelamento.data_vencimento,
        encontrado.linha_parcelamento.valor_historico,
        encontrado.linha_parcelamento.valor_debito,
    ]


def _linha_parcelamento_nao_encontrado_excel(nao_encontrado: DaeParceladoNaoEncontrado) -> list:
    return [
        nao_encontrado.arquivo_origem,
        nao_encontrado.codigo_receita,
        nao_encontrado.referencia,
        nao_encontrado.valor_principal,
        nao_encontrado.status,
    ]


def _ano_mes_parcelamento_encontrado(encontrado: DaeParceladoEncontrado) -> tuple[int, int] | None:
    ano, mes = encontrado.linha_parcelamento.ano_ocorrencia, encontrado.linha_parcelamento.mes_ocorrencia
    return (ano, mes) if ano is not None and mes is not None else None


def gerar_relatorio_parcelamento_excel(
    encontrados: list[DaeParceladoEncontrado],
    nao_encontrados: list[DaeParceladoNaoEncontrado],
    caminho: Path,
) -> Path:
    workbook = Workbook()

    aba_encontrados = workbook.active
    aba_encontrados.title = "Encontrados"
    grupos_encontrados = _agrupar_por_ano_mes(encontrados, _ano_mes_parcelamento_encontrado)
    _escrever_planilha_por_ano_mes(
        aba_encontrados, COLUNAS_PARCELAMENTO_ENCONTRADOS, grupos_encontrados, _linha_parcelamento_encontrado_excel
    )

    aba_nao_encontrados = workbook.create_sheet("Não Encontrados")
    grupos_nao_encontrados = _agrupar_por_ano_mes(nao_encontrados, lambda n: _ano_mes_da_referencia(n.referencia))
    _escrever_planilha_por_ano_mes(
        aba_nao_encontrados,
        COLUNAS_PARCELAMENTO_NAO_ENCONTRADOS,
        grupos_nao_encontrados,
        _linha_parcelamento_nao_encontrado_excel,
    )

    workbook.save(str(caminho))
    return caminho


def gerar_relatorio_parcelamento_pdf(
    encontrados: list[DaeParceladoEncontrado],
    nao_encontrados: list[DaeParceladoNaoEncontrado],
    caminho: Path,
) -> Path:
    documento = SimpleDocTemplate(
        str(caminho),
        pagesize=landscape(A4),
        title="AuditaDAE - Relatorio de Parcelamento",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph("AuditaDAE — Relatório de Parcelamento", estilos["Title"]),
        Paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Spacer(1, 12),
    ]

    elementos.append(Paragraph("DAEs Encontrados no Parcelamento", estilos["Heading2"]))
    grupos_encontrados = _agrupar_por_ano_mes(encontrados, _ano_mes_parcelamento_encontrado)
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_PARCELAMENTO_ENCONTRADOS,
            grupos_encontrados,
            lambda encontrado: [
                encontrado.dae_arquivo_origem,
                encontrado.linha_parcelamento.paf or "",
                encontrado.linha_parcelamento.data_ocorrencia or "",
                encontrado.linha_parcelamento.data_vencimento or "",
                f"{encontrado.linha_parcelamento.valor_historico:.2f}"
                if encontrado.linha_parcelamento.valor_historico is not None
                else "",
                f"{encontrado.linha_parcelamento.valor_debito:.2f}"
                if encontrado.linha_parcelamento.valor_debito is not None
                else "",
            ],
            "#1f6f43",
            "Heading3",
            "Heading4",
        )
    )
    elementos.append(Spacer(1, 12))

    elementos.append(Paragraph("DAEs Não Encontrados no Relatório de Parcelamentos", estilos["Heading2"]))
    grupos_nao_encontrados = _agrupar_por_ano_mes(nao_encontrados, lambda n: _ano_mes_da_referencia(n.referencia))
    elementos.extend(
        _elementos_por_ano_mes(
            estilos,
            COLUNAS_PARCELAMENTO_NAO_ENCONTRADOS,
            grupos_nao_encontrados,
            lambda nao_encontrado: [
                nao_encontrado.arquivo_origem,
                nao_encontrado.codigo_receita or "",
                nao_encontrado.referencia or "",
                f"{nao_encontrado.valor_principal:.2f}" if nao_encontrado.valor_principal is not None else "",
                nao_encontrado.status,
            ],
            "#8f1f1f",
            "Heading3",
            "Heading4",
        )
    )

    documento.build(elementos)
    return caminho
