import openpyxl

ROTULO_NOTA_FISCAL = "nota fiscal"
ROTULO_DESCRICAO = "descri"
COLUNAS_PAGAMENTO_DAE = {
    "nosso número",
    "pagamento",
    "referência",
    "receita",
    "valor principal",
    "valor total",
}

MAX_LINHAS_INSPECIONADAS = 30


def _normalizar(valor) -> str:
    return str(valor).strip().lower() if valor is not None else ""


def classificar_abas(caminho: str) -> dict[str, list[str]]:
    """Varre todas as abas da planilha e classifica cada uma pela estrutura de
    colunas encontrada nas primeiras linhas (não pelo nome da aba), para
    reconhecer o mesmo tipo de dado em planilhas de outros clientes que usem
    nomes de aba diferentes."""
    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        nomes_abas = workbook.sheetnames
    finally:
        workbook.close()

    classificacao: dict[str, list[str]] = {"aquisicao": [], "apuracao": [], "pagamento_dae": []}

    for nome in nomes_abas:
        workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        try:
            planilha = workbook[nome]
            linhas = list(planilha.iter_rows(values_only=True, max_row=MAX_LINHAS_INSPECIONADAS))
        finally:
            workbook.close()

        for linha in linhas:
            valores_normalizados = [_normalizar(c) for c in linha]
            if any(ROTULO_NOTA_FISCAL in v for v in valores_normalizados):
                classificacao["aquisicao"].append(nome)
                break
            if valores_normalizados and valores_normalizados[0].startswith(ROTULO_DESCRICAO):
                classificacao["apuracao"].append(nome)
                break
            if COLUNAS_PAGAMENTO_DAE.issubset(set(valores_normalizados)):
                classificacao["pagamento_dae"].append(nome)
                break

    return classificacao
