import openpyxl

from src.models.dae_models import LinhaRelatorio

NOME_ABA_NF_AQUISICAO = "NF_Aquisicao"
ROTULO_COLUNA_PAI_NOTA_FISCAL = "Nota Fiscal"
SUBROTULO_COLUNA_NOTA = "Nº"


def eh_mapa_aquisicao(caminho: str) -> bool:
    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        return NOME_ABA_NF_AQUISICAO in workbook.sheetnames
    finally:
        workbook.close()


def _localizar_coluna_nota_fiscal(cabecalho_pai: tuple, cabecalho_filho: tuple) -> int:
    for indice, valor in enumerate(cabecalho_pai):
        if not valor or ROTULO_COLUNA_PAI_NOTA_FISCAL not in str(valor):
            continue
        col_inicio = indice
        col_fim = col_inicio
        while col_fim + 1 < len(cabecalho_pai) and cabecalho_pai[col_fim + 1] is None:
            col_fim += 1
        for sub_indice in range(col_inicio, col_fim + 1):
            if sub_indice < len(cabecalho_filho) and cabecalho_filho[sub_indice] and SUBROTULO_COLUNA_NOTA in str(
                cabecalho_filho[sub_indice]
            ):
                return sub_indice
    raise ValueError(
        f"Não foi possível localizar a coluna de número da Nota Fiscal na aba '{NOME_ABA_NF_AQUISICAO}'."
    )


def extrair_notas_mapa_aquisicao(caminho: str, aba: str = NOME_ABA_NF_AQUISICAO) -> list[LinhaRelatorio]:
    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = workbook[aba]
        linhas = list(planilha.iter_rows(values_only=True))
    finally:
        workbook.close()

    indice_cabecalho_pai = next(
        i for i, linha in enumerate(linhas) if any(c and ROTULO_COLUNA_PAI_NOTA_FISCAL in str(c) for c in linha)
    )
    indice_cabecalho_filho = indice_cabecalho_pai + 1
    cabecalho_pai = linhas[indice_cabecalho_pai]
    cabecalho_filho = linhas[indice_cabecalho_filho]

    col_nota = _localizar_coluna_nota_fiscal(cabecalho_pai, cabecalho_filho)
    col_emissao = next((i for i, v in enumerate(cabecalho_filho) if v and "Emiss" in str(v)), 0)
    col_cnpj_emitente = next((i for i, v in enumerate(cabecalho_filho) if v and "CNPJ" in str(v)), None)

    notas_vistas: dict[str, LinhaRelatorio] = {}
    for linha in linhas[indice_cabecalho_filho + 1 :]:
        if col_nota >= len(linha) or linha[col_nota] is None:
            continue
        numero_nf = str(linha[col_nota])
        if numero_nf in notas_vistas:
            continue

        dados_originais = {}
        if col_emissao < len(linha) and linha[col_emissao] is not None:
            dados_originais["Data Emissão"] = str(linha[col_emissao])
        if col_cnpj_emitente is not None and col_cnpj_emitente < len(linha) and linha[col_cnpj_emitente] is not None:
            dados_originais["CNPJ Emitente"] = str(linha[col_cnpj_emitente])

        notas_vistas[numero_nf] = LinhaRelatorio(numero_nf=numero_nf, dados_originais=dados_originais)

    return list(notas_vistas.values())
