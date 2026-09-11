from datetime import datetime

import openpyxl

from src.models.dae_models import LinhaIcmsAt

NOME_ABA_ICMS_AT = "ICMS_AT"


def eh_icms_at_disponivel(caminho: str) -> bool:
    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        return NOME_ABA_ICMS_AT in workbook.sheetnames
    finally:
        workbook.close()


def _valor(linha: tuple, posicao: int) -> float | None:
    if posicao >= len(linha) or linha[posicao] is None:
        return None
    return float(linha[posicao])


def extrair_icms_at(caminho: str, aba: str = NOME_ABA_ICMS_AT) -> list[LinhaIcmsAt]:
    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = workbook[aba]
        linhas = list(planilha.iter_rows(values_only=True))
    finally:
        workbook.close()

    resultado: list[LinhaIcmsAt] = []
    indice = 0
    while indice < len(linhas):
        linha = linhas[indice]
        primeiro = str(linha[0]).strip() if linha and linha[0] else ""
        if primeiro.lower().startswith("descri") and indice + 4 < len(linhas):
            meses = linha[1:13]
            linha_apurado = linhas[indice + 2]
            linha_pago = linhas[indice + 3]
            linha_recolher = linhas[indice + 4]
            for posicao, data_mes in enumerate(meses, start=1):
                if not isinstance(data_mes, datetime):
                    continue
                resultado.append(
                    LinhaIcmsAt(
                        ano=data_mes.year,
                        mes=data_mes.month,
                        valor_apurado=_valor(linha_apurado, posicao),
                        valor_pago=_valor(linha_pago, posicao),
                        valor_a_recolher=_valor(linha_recolher, posicao),
                    )
                )
            indice += 5
        else:
            indice += 1
    return resultado
