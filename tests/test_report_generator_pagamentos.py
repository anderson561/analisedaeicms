import openpyxl
import pdfplumber

from src.models.dae_models import DaePagamentoNaoLocalizado, LinhaPagamentoDae, PagamentoConfirmado
from src.reports.report_generator import gerar_relatorio_pagamentos_excel, gerar_relatorio_pagamentos_pdf


def test_gerar_relatorio_pagamentos_excel_contem_as_duas_abas_agrupadas_por_ano_e_mes(tmp_path):
    confirmados = [
        PagamentoConfirmado(
            dae_arquivo_origem="Macedo DAE 1145 09-2022.pdf",
            linha_pagamento=LinhaPagamentoDae(
                nosso_numero="2113874976",
                data_pagamento="25/02/2022",
                referencia_bruta="1/2022",
                mes_referencia=1,
                ano_referencia=2022,
                codigo_receita="1145",
                descricao_receita="ICMS ANTECIPAÇÃO TRIBUTÁRIA",
                valor_principal=852.93,
                valor_total=852.93,
            ),
        ),
        PagamentoConfirmado(
            dae_arquivo_origem="Macedo DAE 1145 09-2022 (2).pdf",
            linha_pagamento=LinhaPagamentoDae(
                nosso_numero="2113900002",
                data_pagamento="20/10/2022",
                referencia_bruta="9/2022",
                mes_referencia=9,
                ano_referencia=2022,
                codigo_receita="1145",
                descricao_receita="ICMS ANTECIPAÇÃO TRIBUTÁRIA",
                valor_principal=610.00,
                valor_total=610.00,
            ),
        ),
        PagamentoConfirmado(
            dae_arquivo_origem="Macedo DAE 1145 09-2021.pdf",
            linha_pagamento=LinhaPagamentoDae(
                nosso_numero="2100000001",
                data_pagamento="10/02/2021",
                referencia_bruta="1/2021",
                mes_referencia=1,
                ano_referencia=2021,
                codigo_receita="1145",
                descricao_receita="ICMS ANTECIPAÇÃO TRIBUTÁRIA",
                valor_principal=500.00,
                valor_total=500.00,
            ),
        ),
    ]
    nao_localizados = [
        DaePagamentoNaoLocalizado(
            arquivo_origem="Macedo DAE 2175 05-2022.pdf",
            codigo_receita="2175",
            referencia="05/2022",
            valor_principal=784.01,
        ),
        DaePagamentoNaoLocalizado(
            arquivo_origem="Macedo DAE 1145 09-2022.pdf",
            codigo_receita="1145",
            referencia="09/2022",
            valor_principal=999.99,
        ),
        DaePagamentoNaoLocalizado(
            arquivo_origem="Macedo DAE 2175 05-2021.pdf",
            codigo_receita="2175",
            referencia="05/2021",
            valor_principal=100.00,
        ),
    ]
    caminho = tmp_path / "pagamentos.xlsx"

    resultado = gerar_relatorio_pagamentos_excel(confirmados, nao_localizados, caminho)

    assert resultado == caminho
    assert caminho.exists()

    workbook = openpyxl.load_workbook(caminho)
    assert workbook.sheetnames == ["Confirmados", "Não Localizados"]

    aba_confirmados = workbook["Confirmados"]
    valores_confirmados = [[c.value for c in linha] for linha in aba_confirmados.iter_rows()]
    assert valores_confirmados[0][0] == "Ano 2021"
    assert valores_confirmados[1][0] == "Mês 01"
    assert valores_confirmados[2][0] == "Nosso Número"
    assert valores_confirmados[3][0] == "2100000001"
    assert valores_confirmados[4][0] is None
    assert valores_confirmados[5][0] == "Ano 2022"
    assert valores_confirmados[6][0] == "Mês 01"
    assert valores_confirmados[8][0] == "2113874976"
    linha_2022_01 = valores_confirmados[8]
    assert linha_2022_01[1] == "25/02/2022"
    assert linha_2022_01[2] == "1/2022"
    assert linha_2022_01[3] == "1145 - ICMS ANTECIPAÇÃO TRIBUTÁRIA"
    assert linha_2022_01[4] == 852.93
    assert linha_2022_01[5] == 852.93
    assert valores_confirmados[10][0] == "Mês 09"
    assert valores_confirmados[12][0] == "2113900002"

    aba_nao_localizados = workbook["Não Localizados"]
    valores_nao_localizados = [[c.value for c in linha] for linha in aba_nao_localizados.iter_rows()]
    assert valores_nao_localizados[0][0] == "Ano 2021"
    assert valores_nao_localizados[1][0] == "Mês 05"
    assert valores_nao_localizados[3][0] == "Macedo DAE 2175 05-2021.pdf"
    assert valores_nao_localizados[5][0] == "Ano 2022"
    assert valores_nao_localizados[6][0] == "Mês 05"
    linha_nao_localizados_2022_05 = valores_nao_localizados[8]
    assert linha_nao_localizados_2022_05[0] == "Macedo DAE 2175 05-2022.pdf"
    assert linha_nao_localizados_2022_05[1] == "2175"
    assert linha_nao_localizados_2022_05[2] == "05/2022"
    assert linha_nao_localizados_2022_05[3] == 784.01
    assert linha_nao_localizados_2022_05[4] == "Não Localizada no Relatório de Pagamentos"
    assert valores_nao_localizados[10][0] == "Mês 09"
    assert valores_nao_localizados[12][0] == "Macedo DAE 1145 09-2022.pdf"


def test_gerar_relatorio_pagamentos_excel_listas_vazias(tmp_path):
    caminho = tmp_path / "pagamentos_vazio.xlsx"

    resultado = gerar_relatorio_pagamentos_excel([], [], caminho)

    assert resultado.exists()
    workbook = openpyxl.load_workbook(caminho)
    assert workbook.sheetnames == ["Confirmados", "Não Localizados"]


def test_gerar_relatorio_pagamentos_pdf_contem_as_duas_secoes_agrupadas_por_ano_e_mes(tmp_path):
    confirmados = [
        PagamentoConfirmado(
            dae_arquivo_origem="Macedo DAE 2175 05-2022.pdf",
            linha_pagamento=LinhaPagamentoDae(
                nosso_numero="2117979376",
                data_pagamento="27/06/2022",
                referencia_bruta="5/2022",
                mes_referencia=5,
                ano_referencia=2022,
                codigo_receita="2175",
                descricao_receita="ICMS - ANTECIPACAO PARCIAL",
                valor_principal=784.01,
                valor_total=784.01,
            ),
        ),
        PagamentoConfirmado(
            dae_arquivo_origem="Macedo DAE 1145 09-2022.pdf",
            linha_pagamento=LinhaPagamentoDae(
                nosso_numero="2117979999",
                data_pagamento="20/10/2022",
                referencia_bruta="9/2022",
                mes_referencia=9,
                ano_referencia=2022,
                codigo_receita="1145",
                descricao_receita="ICMS ANTECIPAÇÃO TRIBUTÁRIA",
                valor_principal=610.00,
                valor_total=610.00,
            ),
        ),
        PagamentoConfirmado(
            dae_arquivo_origem="Macedo DAE 2175 05-2021.pdf",
            linha_pagamento=LinhaPagamentoDae(
                nosso_numero="2100000002",
                data_pagamento="27/06/2021",
                referencia_bruta="5/2021",
                mes_referencia=5,
                ano_referencia=2021,
                codigo_receita="2175",
                descricao_receita="ICMS - ANTECIPACAO PARCIAL",
                valor_principal=200.00,
                valor_total=200.00,
            ),
        ),
    ]
    nao_localizados = [
        DaePagamentoNaoLocalizado(
            arquivo_origem="Macedo DAE 1145 09-2022.pdf",
            codigo_receita="1145",
            referencia="09/2022",
            valor_principal=4998.63,
        ),
        DaePagamentoNaoLocalizado(
            arquivo_origem="Macedo DAE 1145 09-2021.pdf",
            codigo_receita="1145",
            referencia="09/2021",
            valor_principal=1000.00,
        ),
    ]
    caminho = tmp_path / "pagamentos.pdf"

    resultado = gerar_relatorio_pagamentos_pdf(confirmados, nao_localizados, caminho)

    assert resultado == caminho
    assert caminho.exists()

    with pdfplumber.open(caminho) as pdf:
        texto = "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)

    assert "Relatório de Pagamentos" in texto
    assert "Pagamentos Confirmados" in texto
    assert "2117979376" in texto
    assert "784.01" in texto
    assert "Não Localizados" in texto
    assert "Macedo DAE 1145 09-2022.pdf" in texto
    assert "4998.63" in texto

    assert texto.count("Ano 2021") == 2
    assert texto.count("Ano 2022") == 2
    secao_confirmados = texto.index("Pagamentos Confirmados")
    secao_nao_localizados = texto.index("Não Localizados no Relatório")
    ano_2021_confirmados = texto.index("Ano 2021", secao_confirmados)
    ano_2022_confirmados = texto.index("Ano 2022", secao_confirmados)
    assert secao_confirmados < ano_2021_confirmados < ano_2022_confirmados < secao_nao_localizados

    mes_05_confirmados = texto.index("Mês 05", ano_2022_confirmados)
    mes_09_confirmados = texto.index("Mês 09", ano_2022_confirmados)
    assert ano_2022_confirmados < mes_05_confirmados < mes_09_confirmados < secao_nao_localizados


def test_gerar_relatorio_pagamentos_pdf_listas_vazias(tmp_path):
    caminho = tmp_path / "pagamentos_vazio.pdf"

    resultado = gerar_relatorio_pagamentos_pdf([], [], caminho)

    assert resultado.exists()
