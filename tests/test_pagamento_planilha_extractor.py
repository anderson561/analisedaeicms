from datetime import datetime

import openpyxl

from src.extractors.pagamento_planilha_extractor import extrair_pagamentos_e_parcelamento_de_planilha


def _construir_planilha_pagamento_dae(caminho, com_preambulo: bool):
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Arrecadacao"

    if com_preambulo:
        aba.append(["MACEDO COMERCIAL DE CALCADOS LTDA"])
        aba.append(["Endereco qualquer"])
        aba.append([])
        aba.append(["Relação de DAEs"])

    aba.append(["Nosso Número", "Pagamento", "Referência", "Receita", "Valor Principal", "Valor Total"])
    aba.append([1001, "26/12/2025", datetime(2025, 11, 1), "1.145 - ICMS ANTECIPAÇÃO TRIBUTÁRIA", 5068.29, 5068.29])
    aba.append([1002, "19/01/2026", "015/023", "1.802 - ICMS PARCELAMENTO", 1494.50, 1926.11])
    workbook.save(caminho)


def test_extrai_linha_com_referencia_data_como_pagamento(tmp_path):
    caminho = tmp_path / "arrecadacao.xlsx"
    _construir_planilha_pagamento_dae(caminho, com_preambulo=True)

    pagamentos, parcelamentos = extrair_pagamentos_e_parcelamento_de_planilha(str(caminho), ["Arrecadacao"])

    assert len(pagamentos) == 1
    linha = pagamentos[0]
    assert linha.nosso_numero == "1001"
    assert linha.mes_referencia == 11
    assert linha.ano_referencia == 2025
    assert linha.codigo_receita == "1.145"
    assert linha.descricao_receita == "ICMS ANTECIPAÇÃO TRIBUTÁRIA"
    assert linha.valor_principal == 5068.29
    assert linha.valor_total == 5068.29


def test_extrai_linha_com_referencia_parcela_como_parcelamento(tmp_path):
    caminho = tmp_path / "arrecadacao.xlsx"
    _construir_planilha_pagamento_dae(caminho, com_preambulo=True)

    pagamentos, parcelamentos = extrair_pagamentos_e_parcelamento_de_planilha(str(caminho), ["Arrecadacao"])

    assert len(parcelamentos) == 1
    linha = parcelamentos[0]
    assert linha.arquivo_origem == "Arrecadacao"
    assert linha.data_vencimento == "19/01/2026"
    assert linha.valor_historico == 1494.50
    assert linha.valor_debito == 1926.11
    assert linha.mes_ocorrencia is None
    assert linha.ano_ocorrencia is None


def test_localiza_cabecalho_mesmo_com_linhas_de_preambulo(tmp_path):
    caminho = tmp_path / "arrecadacao.xlsx"
    _construir_planilha_pagamento_dae(caminho, com_preambulo=True)

    pagamentos, parcelamentos = extrair_pagamentos_e_parcelamento_de_planilha(str(caminho), ["Arrecadacao"])

    assert len(pagamentos) + len(parcelamentos) == 2


def test_multiplas_abas_sao_combinadas(tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    workbook = openpyxl.Workbook()
    aba1 = workbook.active
    aba1.title = "Arrecadacao"
    aba1.append(["Nosso Número", "Pagamento", "Referência", "Receita", "Valor Principal", "Valor Total"])
    aba1.append([1001, "26/12/2025", datetime(2025, 11, 1), "1.145 - ICMS", 100.0, 100.0])

    aba2 = workbook.create_sheet("Parcelamento")
    aba2.append(["Nosso Número", "Pagamento", "Referência", "Receita", "Valor Principal", "Valor Total"])
    aba2.append([2002, "10/01/2026", "001/012", "1.802 - ICMS PARCELAMENTO", 200.0, 250.0])

    workbook.save(caminho)

    pagamentos, parcelamentos = extrair_pagamentos_e_parcelamento_de_planilha(
        str(caminho), ["Arrecadacao", "Parcelamento"]
    )

    assert len(pagamentos) == 1
    assert len(parcelamentos) == 1
