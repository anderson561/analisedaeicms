import openpyxl

from src.extractors.classificador_planilha import classificar_abas


def test_classifica_aba_de_aquisicao_com_nome_diferente(tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Compras do Mes"
    aba.append(["Data", "Data", "Nota Fiscal", None, None, "Chave", "Emitente"])
    aba.append([None, None, "Nº", None, None, None, None])
    workbook.save(caminho)

    classificacao = classificar_abas(str(caminho))

    assert classificacao["aquisicao"] == ["Compras do Mes"]
    assert classificacao["apuracao"] == []
    assert classificacao["pagamento_dae"] == []


def test_classifica_aba_de_apuracao_com_nome_diferente(tmp_path):
    from datetime import datetime

    caminho = tmp_path / "planilha.xlsx"
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Resumo Fiscal"
    aba.append(["Descrição"] + [datetime(2021, mes, 1) for mes in range(1, 13)])
    workbook.save(caminho)

    classificacao = classificar_abas(str(caminho))

    assert classificacao["apuracao"] == ["Resumo Fiscal"]
    assert classificacao["aquisicao"] == []
    assert classificacao["pagamento_dae"] == []


def test_classifica_aba_de_pagamento_dae_com_nome_diferente(tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Historico de DAEs"
    aba.append(["Título qualquer"])
    aba.append([])
    aba.append(["Nosso Número", "Pagamento", "Referência", "Receita", "Valor Principal", "Valor Total"])
    aba.append([123, "01/01/2024", "01/2024", "1.145 - ICMS", 100.0, 100.0])
    workbook.save(caminho)

    classificacao = classificar_abas(str(caminho))

    assert classificacao["pagamento_dae"] == ["Historico de DAEs"]
    assert classificacao["aquisicao"] == []
    assert classificacao["apuracao"] == []


def test_classifica_multiplas_abas_do_mesmo_tipo(tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    workbook = openpyxl.Workbook()
    aba1 = workbook.active
    aba1.title = "Arrecadacao"
    aba1.append(["Nosso Número", "Pagamento", "Referência", "Receita", "Valor Principal", "Valor Total"])

    aba2 = workbook.create_sheet("Parcelamento")
    aba2.append(["Nosso Número", "Pagamento", "Referência", "Receita", "Valor Principal", "Valor Total"])

    workbook.save(caminho)

    classificacao = classificar_abas(str(caminho))

    assert set(classificacao["pagamento_dae"]) == {"Arrecadacao", "Parcelamento"}


def test_aba_sem_padrao_reconhecido_e_ignorada(tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Notas Explicativas"
    aba.append(["Algum texto qualquer", "Outra coluna"])
    workbook.save(caminho)

    classificacao = classificar_abas(str(caminho))

    assert classificacao == {"aquisicao": [], "apuracao": [], "pagamento_dae": []}
