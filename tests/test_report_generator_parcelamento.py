import openpyxl
import pdfplumber

from src.models.dae_models import DaeParceladoEncontrado, DaeParceladoNaoEncontrado, LinhaParcelamento
from src.reports.report_generator import gerar_relatorio_parcelamento_excel, gerar_relatorio_parcelamento_pdf


def test_gerar_relatorio_parcelamento_excel_contem_as_duas_abas_agrupadas_por_ano_e_mes(tmp_path):
    encontrados = [
        DaeParceladoEncontrado(
            dae_arquivo_origem="Macedo DAE 07.02.004 07-2019.pdf",
            linha_parcelamento=LinhaParcelamento(
                paf="810000.7810/24-5",
                arquivo_origem="2022.pdf",
                data_ocorrencia="31/07/2019",
                data_vencimento="25/08/2019",
                mes_ocorrencia=7,
                ano_ocorrencia=2019,
                valor_historico=1009.63,
                valor_debito=1009.63,
            ),
        ),
        DaeParceladoEncontrado(
            dae_arquivo_origem="Macedo DAE 07.02.004 09-2022.pdf",
            linha_parcelamento=LinhaParcelamento(
                paf="810000.7810/24-5",
                arquivo_origem="2022.pdf",
                data_ocorrencia="30/09/2022",
                data_vencimento="25/10/2022",
                mes_ocorrencia=9,
                ano_ocorrencia=2022,
                valor_historico=384.35,
                valor_debito=384.35,
            ),
        ),
    ]
    nao_encontrados = [
        DaeParceladoNaoEncontrado(
            arquivo_origem="Macedo DAE 1145 09-2022.pdf",
            codigo_receita="1145",
            referencia="09/2022",
            valor_principal=999.99,
        ),
    ]
    caminho = tmp_path / "parcelamento.xlsx"

    resultado = gerar_relatorio_parcelamento_excel(encontrados, nao_encontrados, caminho)

    assert resultado == caminho
    assert caminho.exists()

    workbook = openpyxl.load_workbook(caminho)
    assert workbook.sheetnames == ["Encontrados", "Não Encontrados"]

    aba_encontrados = workbook["Encontrados"]
    valores_encontrados = [[c.value for c in linha] for linha in aba_encontrados.iter_rows()]
    assert valores_encontrados[0][0] == "Ano 2019"
    assert valores_encontrados[1][0] == "Mês 07"
    assert valores_encontrados[2][0] == "Arquivo DAE"
    assert valores_encontrados[3][0] == "Macedo DAE 07.02.004 07-2019.pdf"
    linha_2019 = valores_encontrados[3]
    assert linha_2019[1] == "810000.7810/24-5"
    assert linha_2019[2] == "31/07/2019"
    assert linha_2019[3] == "25/08/2019"
    assert linha_2019[4] == 1009.63
    assert linha_2019[5] == 1009.63
    assert valores_encontrados[5][0] == "Ano 2022"
    assert valores_encontrados[6][0] == "Mês 09"
    assert valores_encontrados[8][0] == "Macedo DAE 07.02.004 09-2022.pdf"

    aba_nao_encontrados = workbook["Não Encontrados"]
    valores_nao_encontrados = [[c.value for c in linha] for linha in aba_nao_encontrados.iter_rows()]
    assert valores_nao_encontrados[0][0] == "Ano 2022"
    assert valores_nao_encontrados[1][0] == "Mês 09"
    linha_nao_encontrada = valores_nao_encontrados[3]
    assert linha_nao_encontrada[0] == "Macedo DAE 1145 09-2022.pdf"
    assert linha_nao_encontrada[1] == "1145"
    assert linha_nao_encontrada[2] == "09/2022"
    assert linha_nao_encontrada[3] == 999.99
    assert linha_nao_encontrada[4] == "Não Localizado no Relatório de Parcelamentos"


def test_gerar_relatorio_parcelamento_excel_listas_vazias(tmp_path):
    caminho = tmp_path / "parcelamento_vazio.xlsx"

    resultado = gerar_relatorio_parcelamento_excel([], [], caminho)

    assert resultado.exists()
    workbook = openpyxl.load_workbook(caminho)
    assert workbook.sheetnames == ["Encontrados", "Não Encontrados"]


def test_gerar_relatorio_parcelamento_pdf_contem_as_duas_secoes_agrupadas_por_ano_e_mes(tmp_path):
    encontrados = [
        DaeParceladoEncontrado(
            dae_arquivo_origem="Macedo DAE 07.02.004 07-2019.pdf",
            linha_parcelamento=LinhaParcelamento(
                paf="810000.7810/24-5",
                arquivo_origem="2022.pdf",
                data_ocorrencia="31/07/2019",
                data_vencimento="25/08/2019",
                mes_ocorrencia=7,
                ano_ocorrencia=2019,
                valor_historico=1009.63,
                valor_debito=1009.63,
            ),
        ),
    ]
    nao_encontrados = [
        DaeParceladoNaoEncontrado(
            arquivo_origem="Macedo DAE 1145 09-2022.pdf",
            codigo_receita="1145",
            referencia="09/2022",
            valor_principal=999.99,
        ),
    ]
    caminho = tmp_path / "parcelamento.pdf"

    resultado = gerar_relatorio_parcelamento_pdf(encontrados, nao_encontrados, caminho)

    assert resultado == caminho
    assert caminho.exists()

    with pdfplumber.open(caminho) as pdf:
        texto = "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)

    assert "Relatório de Parcelamento" in texto
    assert "DAEs Encontrados" in texto
    assert "Macedo DAE 07.02.004 07-2019.pdf" in texto
    assert "810000.7810/24-5" in texto
    assert "1009.63" in texto
    assert "DAEs Não Encontrados" in texto
    assert "Macedo DAE 1145 09-2022.pdf" in texto
    assert "999.99" in texto
    assert "Ano 2019" in texto
    assert "Mês 07" in texto


def test_gerar_relatorio_parcelamento_pdf_listas_vazias(tmp_path):
    caminho = tmp_path / "parcelamento_vazio.pdf"

    resultado = gerar_relatorio_parcelamento_pdf([], [], caminho)

    assert resultado.exists()
